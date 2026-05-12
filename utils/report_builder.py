"""
Enhanced report generator.
- Renders Plotly charts to PNG (embedded in HTML / PDF / DOCX)
- Calls LLM to generate detailed section-by-section analysis
- Supports Markdown / HTML / PDF / DOCX output formats
"""
import base64
import io
from typing import Dict, List, Optional
import pandas as pd

from utils.charts import (
    plot_trend, plot_dupont_waterfall, plot_dupont_decomposition,
    plot_peer_comparison, plot_radar
)
from utils.ratios import is_percentage_metric


# ============================================================
# Plotly figure -> PNG (base64)
# ============================================================
def fig_to_png_base64(fig, width: int = 900, height: int = 450) -> Optional[str]:
    """Convert a Plotly figure to base64 PNG string. Returns None on failure."""
    try:
        png_bytes = fig.to_image(format="png", width=width, height=height, scale=2)
        return base64.b64encode(png_bytes).decode("utf-8")
    except Exception:
        return None


def fig_to_png_bytes(fig, width: int = 900, height: int = 450) -> Optional[bytes]:
    """Plotly figure -> raw PNG bytes (for docx embedding)."""
    try:
        return fig.to_image(format="png", width=width, height=height, scale=2)
    except Exception:
        return None


# ============================================================
# Collect all charts
# ============================================================
def collect_all_charts(ratios: Dict, dupont: Dict, trend_df: pd.DataFrame,
                        compare_df: pd.DataFrame, year: int) -> List[Dict]:
    """
    Build all chart objects. Returns:
    [{"id": "...", "title": "...", "category": "...",
      "section_anchor": "...", "fig": <plotly>}, ...]

    section_anchor maps charts to their report sections.
    """
    charts = []

    # 1. DuPont waterfall -> "dupont" section
    charts.append({
        "id": "dupont_waterfall",
        "title": f"{year} DuPont Three-Factor Decomposition",
        "category": "DuPont",
        "section_anchor": "dupont",
        "fig": plot_dupont_waterfall(dupont, year),
    })

    # 2. DuPont multi-year trend -> "dupont" section
    if not trend_df.empty:
        dupont_history = {}
        for col in trend_df.columns:
            yr = trend_df[col].to_dict()
            dupont_history[col] = {
                "Net Margin": yr.get("Net Margin"),
                "Asset Turnover": yr.get("Asset Turnover"),
                "Equity Multiplier": yr.get("Equity Multiplier"),
                "ROE (DuPont)": yr.get("ROE (Return on Equity)"),
            }
        charts.append({
            "id": "dupont_trend",
            "title": "DuPont Three-Factor Trends",
            "category": "DuPont",
            "section_anchor": "dupont",
            "fig": plot_dupont_decomposition(dupont_history),
        })

    # 3. Profitability trends -> "profitability"
    if not trend_df.empty:
        charts.append({
            "id": "trend_profit",
            "title": "Profitability Trends (Gross / Operating / Net Margin)",
            "category": "Trend",
            "section_anchor": "profitability",
            "fig": plot_trend(trend_df,
                ["Gross Margin", "Operating Margin", "Net Margin"],
                title="Profitability Trends"),
        })
        # 4. Return trends -> "trend"
        charts.append({
            "id": "trend_return",
            "title": "Return Trends (ROE / ROA)",
            "category": "Trend",
            "section_anchor": "trend",
            "fig": plot_trend(trend_df,
                ["ROE (Return on Equity)", "ROA (Return on Assets)"],
                title="Return Trends"),
        })
        # 5. Solvency trends -> "solvency"
        charts.append({
            "id": "trend_solvency",
            "title": "Solvency Trends (Current / Quick / Debt to Assets)",
            "category": "Trend",
            "section_anchor": "solvency",
            "fig": plot_trend(trend_df,
                ["Current Ratio", "Quick Ratio", "Debt to Assets"],
                title="Solvency Trends"),
        })

    # 6-9. Peer comparison bars + radar -> "peer" section
    if not compare_df.empty and len(compare_df.columns) >= 2:
        for m in ["Net Margin", "ROE (Return on Equity)",
                  "ROA (Return on Assets)", "Debt to Assets"]:
            if m in compare_df.index:
                charts.append({
                    "id": f"peer_{m}",
                    "title": f"Peer Comparison: {m}",
                    "category": "Peer",
                    "section_anchor": "peer",
                    "fig": plot_peer_comparison(compare_df, m),
                })

        radar_metrics = ["Net Margin", "ROE (Return on Equity)",
                         "ROA (Return on Assets)", "Asset Turnover",
                         "Gross Margin", "Current Ratio"]
        charts.append({
            "id": "radar",
            "title": "Performance Radar (Target vs. Peer Average)",
            "category": "Peer",
            "section_anchor": "peer",
            "fig": plot_radar(compare_df, radar_metrics),
        })

    return charts


# ============================================================
# LLM section-by-section analysis
# ============================================================
def generate_section_analyses(company_info: Dict, ratios: Dict, dupont: Dict,
                                trend_df: pd.DataFrame, compare_df: pd.DataFrame,
                                year: int, api_key: str,
                                model: str = "gpt-4o-mini") -> Dict[str, str]:
    """
    Use LLM to generate per-section analysis based on actual data.
    Returns a dict with 8 keys: overview, profitability, operating, solvency,
                                 dupont, trend, peer, diagnosis
    """
    from openai import OpenAI
    import json as _json

    def clean_ratios(d):
        return {k: round(v, 4) if v is not None else None
                for k, v in d.items() if not k.startswith("_")}

    data = {
        "company_info": {
            "name": company_info.get("name"),
            "ticker": company_info.get("ticker"),
            "sector": company_info.get("sector"),
            "industry": company_info.get("industry"),
            "country": company_info.get("country"),
            "market_cap": company_info.get("market_cap"),
            "currency": company_info.get("currency"),
            "business_summary": (company_info.get("summary") or "")[:300],
        },
        "fiscal_year": year,
        "all_ratios": clean_ratios(ratios),
        "dupont_breakdown": {k: round(v, 4) if v else None for k, v in dupont.items()},
    }

    if not trend_df.empty:
        trend_dict = {}
        for idx in trend_df.index:
            trend_dict[idx] = {str(c): round(v, 4) if pd.notna(v) else None
                                for c, v in trend_df.loc[idx].items()}
        data["multi_year_trends"] = trend_dict

    if not compare_df.empty:
        peer_dict = {}
        for idx in compare_df.index:
            peer_dict[idx] = {str(c): round(v, 4) if pd.notna(v) else None
                                for c, v in compare_df.loc[idx].items()}
        data["peer_comparison"] = peer_dict

    system_prompt = """You are a senior financial analyst. The user will give you complete financial data for a US-listed company, and you need to write text analysis for **8 sections** of a sectional financial report.

# Critical requirements (read first)
1. The user experience is: **they see charts/data tables first, then read your text analysis**. So your text should frequently reference specific numbers, as if interpreting the chart/table that appeared above.
2. Do NOT say "in the table below" or "as shown in the chart below" - the charts come BEFORE your text. Use phrases like "the table above shows", "the chart above indicates", "as visible in the figure".
3. Never paste Markdown tables in your text (the system already renders the actual tables).
4. Each paragraph should provide **professional insight** grounded in specific numbers, not just restate them.
5. Consider industry context (banks have high leverage normally, tech companies have high gross margins as expected, retail relies on turnover, etc.).
6. When multi-year trend data exists, always discuss "direction" along with the trend.
7. When peer data exists, always give a "relative position" assessment.

# Responsibilities of the 8 sections
- **overview**: Company overview. Based on company info and market cap, briefly assess market position, industry profile, and scale. ~150-250 words.
- **profitability**: Profitability analysis. Based on the profitability chart (gross/operating/net margin trends) and profitability ratios, deeply analyze profit quality, structure, and trend changes. ~250-350 words.
- **operating**: Operating efficiency. Based on asset turnover, inventory turnover, receivables turnover, analyze efficiency given industry context. ~200-300 words.
- **solvency**: Solvency and capital structure. Based on the solvency trend chart (current/quick/debt-to-assets) and ratios, analyze short-term liquidity and long-term solvency. ~200-300 words.
- **dupont**: DuPont analysis. Based on the DuPont waterfall (single year) and DuPont historical trends (how each factor changed), decompose ROE drivers and identify which factor is changing. ~250-350 words.
- **trend**: Multi-year overall trend. Based on the return trend chart (ROE/ROA) and all trend data, discuss the company's trajectory and directional changes. ~200-300 words.
- **peer**: Peer comparison. Based on peer comparison bar charts (Net Margin/ROE/ROA/Debt to Assets) and radar chart, analyze relative position among peers, strengths and weaknesses. If no peer data, write "No peer data provided". ~250-350 words.
- **diagnosis**: Overall diagnosis and investor perspective. Combining all above, use three H3 headers "### Strengths", "### Risks & Weaknesses", "### Key Points for Investors" to give overall judgment. ~300-400 words.

# Output format (strict)
Return a JSON object directly. **Do NOT** wrap in ```json``` code blocks. No preamble or commentary. Structure:
```
{"overview": "text", "profitability": "text", "operating": "text", "solvency": "text", "dupont": "text", "trend": "text", "peer": "text", "diagnosis": "text"}
```

Each value is Markdown text (can use **bold**, - lists, ### H3 headers). Do NOT use # H1 or ## H2 headers (the system handles section titles)."""

    user_prompt = f"""Based on the following data, generate section-by-section financial analysis for **{data['company_info']['name']}** ({data['company_info']['ticker']}) for fiscal year {year}.

```json
{_json.dumps(data, ensure_ascii=False, indent=2)}
```

Return the JSON object directly, starting with {{ and ending with }}."""

    client = OpenAI(api_key=api_key)
    resp = client.chat.completions.create(
        model=model,
        messages=[
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        temperature=0.4,
        max_tokens=4000,
        response_format={"type": "json_object"},
    )
    content = resp.choices[0].message.content or "{}"

    try:
        sections = _json.loads(content)
    except _json.JSONDecodeError:
        sections = {}

    expected_keys = ["overview", "profitability", "operating", "solvency",
                     "dupont", "trend", "peer", "diagnosis"]
    for k in expected_keys:
        if k not in sections:
            sections[k] = ""

    return sections


# Backward-compat: combine sections into one Markdown string
def generate_llm_analysis(company_info: Dict, ratios: Dict, dupont: Dict,
                            trend_df: pd.DataFrame, compare_df: pd.DataFrame,
                            year: int, api_key: str,
                            model: str = "gpt-4o-mini") -> str:
    """Legacy API wrapper: returns combined Markdown text."""
    sections = generate_section_analyses(
        company_info, ratios, dupont, trend_df, compare_df, year, api_key, model
    )
    titles = {
        "overview": "## 1. Company Overview",
        "profitability": "## 2. Profitability Analysis",
        "operating": "## 3. Operating Efficiency",
        "solvency": "## 4. Solvency and Capital Structure",
        "dupont": "## 5. DuPont Analysis: ROE Drivers",
        "trend": "## 6. Multi-Year Trends",
        "peer": "## 7. Peer Comparison",
        "diagnosis": "## 8. Overall Diagnosis and Investor Perspective",
    }
    parts = []
    for key, title in titles.items():
        body = sections.get(key, "").strip()
        if body:
            parts.append(f"{title}\n\n{body}")
    return "\n\n".join(parts)


# ============================================================
# HTML report
# ============================================================
HTML_CSS = """
<style>
  body {
    font-family: 'Helvetica Neue', Arial, sans-serif;
    max-width: 920px;
    margin: 32px auto;
    padding: 24px;
    color: #222;
    line-height: 1.7;
    background: #fff;
  }
  h1 { color: #1f4e79; border-bottom: 3px solid #c00000; padding-bottom: 12px; }
  h2 { color: #1f4e79; border-left: 4px solid #c00000; padding-left: 12px;
        margin-top: 36px; }
  h3 { color: #2c5282; margin-top: 24px; }
  .meta { color: #555; font-size: 14px; margin-bottom: 24px; }
  .chart-block { margin: 24px 0; text-align: center; }
  .chart-block img { max-width: 100%; height: auto;
                       border: 1px solid #e5e7eb; border-radius: 6px; }
  .chart-caption { font-size: 13px; color: #666; margin-top: 6px; font-style: italic; }
  table { width: 100%; border-collapse: collapse; margin: 16px 0; font-size: 14px; }
  th, td { border: 1px solid #ddd; padding: 8px 12px; text-align: left; }
  th { background: #f5f7fa; color: #1f4e79; }
  tr:nth-child(even) { background: #fafbfc; }
  .summary-box { background: #f5f7fa; border-left: 4px solid #1f4e79;
                  padding: 16px 20px; margin: 20px 0; border-radius: 4px; }
  .footer { margin-top: 48px; padding-top: 16px; border-top: 1px solid #eee;
              color: #888; font-size: 12px; text-align: center; }
  ul { padding-left: 24px; }
  li { margin: 6px 0; }
  strong { color: #1f4e79; }
</style>
"""


def md_to_html(md: str) -> str:
    """Lightweight Markdown -> HTML converter (basic elements only)."""
    import re
    lines = md.split("\n")
    out = []
    in_list = False
    for line in lines:
        line = line.rstrip()
        if not line:
            if in_list:
                out.append("</ul>")
                in_list = False
            out.append("")
            continue
        m = re.match(r"^(#{1,4})\s+(.+)", line)
        if m:
            if in_list:
                out.append("</ul>")
                in_list = False
            level = len(m.group(1))
            out.append(f"<h{level}>{m.group(2)}</h{level}>")
            continue
        m_li = re.match(r"^[-*]\s+(.+)", line)
        if m_li:
            if not in_list:
                out.append("<ul>")
                in_list = True
            out.append(f"<li>{_inline_md(m_li.group(1))}</li>")
            continue
        if in_list:
            out.append("</ul>")
            in_list = False
        out.append(f"<p>{_inline_md(line)}</p>")
    if in_list:
        out.append("</ul>")
    return "\n".join(out)


def _inline_md(text: str) -> str:
    """Handle inline **bold** *italic* `code`."""
    import re
    text = re.sub(r"\*\*([^\*]+)\*\*", r"<strong>\1</strong>", text)
    text = re.sub(r"\*([^\*]+)\*", r"<em>\1</em>", text)
    text = re.sub(r"`([^`]+)`", r"<code>\1</code>", text)
    return text


def _format_value(value, metric_name: str) -> str:
    """Format a metric value for display, choosing % or raw based on metric type."""
    if value is None or pd.isna(value):
        return "-"
    is_pct = is_percentage_metric(metric_name)
    if is_pct:
        return f"{value*100:.2f}%"
    if abs(value) > 1e6:
        return f"{value/1e6:.1f}M"
    return f"{value:.3f}"


def _build_ratios_subset_table(ratios: Dict, keys: List[str]) -> str:
    """Build an HTML table for a subset of ratio keys."""
    rows_html = ""
    has_data = False
    for k in keys:
        v = ratios.get(k)
        if v is None and k not in ratios:
            continue
        disp = _format_value(v, k)
        rows_html += f"<tr><td>{k}</td><td><strong>{disp}</strong></td></tr>"
        has_data = True
    if not has_data:
        return ""
    return f"""<table><thead><tr><th>Metric</th><th>Value</th></tr></thead>
<tbody>{rows_html}</tbody></table>"""


def _build_company_info_table(info: Dict) -> str:
    """Company info HTML table."""
    mcap = info.get("market_cap")
    mcap_str = f"${mcap/1e9:.1f}B {info.get('currency', 'USD')}" if mcap else "-"
    rows = [
        ("Company Name", info.get("name", "-")),
        ("Ticker", info.get("ticker", "-")),
        ("Sector", info.get("sector", "-")),
        ("Industry", info.get("industry", "-")),
        ("Country", info.get("country", "-")),
        ("Market Cap", mcap_str),
        ("Reporting Currency", info.get("currency", "-")),
    ]
    rows_html = "".join(f"<tr><td>{k}</td><td>{v}</td></tr>" for k, v in rows)
    return f"<table><tbody>{rows_html}</tbody></table>"


def _build_peer_table(compare_df: pd.DataFrame) -> str:
    """Peer comparison HTML table."""
    if compare_df.empty:
        return ""
    html = "<table><thead><tr><th>Metric</th>"
    for col in compare_df.columns:
        html += f"<th>{col}</th>"
    html += "</tr></thead><tbody>"
    for idx in compare_df.index:
        html += f"<tr><td>{idx}</td>"
        for col in compare_df.columns:
            v = compare_df.loc[idx, col]
            html += f"<td>{_format_value(v, idx)}</td>"
        html += "</tr>"
    html += "</tbody></table>"
    return html


def build_html_report(company_info: Dict, year: int, sections: Dict[str, str],
                       charts: List[Dict], ratios: Dict,
                       compare_df: pd.DataFrame) -> str:
    """
    Build HTML report - rendered as "chart/table -> text" per section.
    sections: 8 sections from generate_section_analyses
    """
    name = company_info.get("name", "Company")
    ticker = company_info.get("ticker", "")
    sector = company_info.get("sector", "-")
    industry = company_info.get("industry", "-")
    currency = company_info.get("currency", "USD")
    mcap = company_info.get("market_cap")
    mcap_str = f"${mcap/1e9:.1f}B {currency}" if mcap else "-"

    # Render charts to base64, indexed by id
    chart_blocks = {}
    for chart in charts:
        b64 = fig_to_png_base64(chart["fig"])
        if b64:
            chart_blocks[chart["id"]] = f"""
<div class="chart-block">
    <img src="data:image/png;base64,{b64}" alt="{chart['title']}" />
    <div class="chart-caption">Figure: {chart['title']}</div>
</div>"""
        else:
            chart_blocks[chart["id"]] = (
                f'<p class="chart-caption">Chart "{chart["title"]}" could not be rendered</p>'
            )

    def render_md(text: str) -> str:
        if not text or not text.strip():
            return "<p><em>(No analysis available for this section.)</em></p>"
        return f'<div class="section-text">{md_to_html(text)}</div>'

    profit_keys = ["Gross Margin", "Operating Margin", "Net Margin",
                   "ROA (Return on Assets)", "ROE (Return on Equity)"]
    operating_keys = ["Asset Turnover", "Inventory Turnover", "Receivables Turnover"]
    solvency_keys = ["Current Ratio", "Quick Ratio", "Debt to Assets",
                     "Equity Multiplier", "Interest Coverage"]

    def chart_ids_for(anchor: str) -> List[str]:
        return [c["id"] for c in charts if c.get("section_anchor") == anchor]

    def insert_charts(ids: List[str]) -> str:
        return "".join(chart_blocks.get(i, "") for i in ids)

    body_html = ""

    # 1. Company Overview
    body_html += "<h2>1. Company Overview</h2>\n"
    body_html += _build_company_info_table(company_info)
    body_html += render_md(sections.get("overview", ""))

    # 2. Profitability
    body_html += "<h2>2. Profitability Analysis</h2>\n"
    body_html += insert_charts(chart_ids_for("profitability"))
    body_html += "<h3>Key Profitability Metrics</h3>"
    body_html += _build_ratios_subset_table(ratios, profit_keys)
    body_html += render_md(sections.get("profitability", ""))

    # 3. Operating Efficiency
    body_html += "<h2>3. Operating Efficiency</h2>\n"
    body_html += "<h3>Efficiency Metrics</h3>"
    body_html += _build_ratios_subset_table(ratios, operating_keys)
    body_html += render_md(sections.get("operating", ""))

    # 4. Solvency
    body_html += "<h2>4. Solvency and Capital Structure</h2>\n"
    body_html += insert_charts(chart_ids_for("solvency"))
    body_html += "<h3>Solvency Metrics</h3>"
    body_html += _build_ratios_subset_table(ratios, solvency_keys)
    body_html += render_md(sections.get("solvency", ""))

    # 5. DuPont
    body_html += "<h2>5. DuPont Analysis: ROE Drivers</h2>\n"
    body_html += insert_charts(chart_ids_for("dupont"))
    body_html += render_md(sections.get("dupont", ""))

    # 6. Multi-Year Trends
    body_html += "<h2>6. Multi-Year Trends</h2>\n"
    body_html += insert_charts(chart_ids_for("trend"))
    body_html += render_md(sections.get("trend", ""))

    # 7. Peer Comparison
    body_html += "<h2>7. Peer Comparison</h2>\n"
    if not compare_df.empty:
        body_html += "<h3>Peer Comparison Table</h3>"
        body_html += _build_peer_table(compare_df)
        body_html += insert_charts(chart_ids_for("peer"))
    body_html += render_md(sections.get("peer", ""))

    # 8. Diagnosis
    body_html += "<h2>8. Overall Diagnosis and Investor Perspective</h2>\n"
    body_html += render_md(sections.get("diagnosis", ""))

    return f"""<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="UTF-8">
<title>{name} ({ticker}) FY{year} Financial Analysis</title>
{HTML_CSS}
</head>
<body>

<h1>{name} ({ticker})</h1>
<p class="meta">
<strong>FY{year} Financial Analysis Report</strong><br>
Sector: {sector} / {industry} &nbsp;|&nbsp; Country: {company_info.get('country', '-')} \
&nbsp;|&nbsp; Market Cap: {mcap_str}
</p>

<div class="summary-box">
This report is auto-generated by the AI Financial Analysis Agent. Each section first \
presents key data (charts and tables), followed by AI-driven analysis built on that data.
</div>

{body_html}

<div class="footer">
This report is generated from public financial data (yfinance / FMP). It is for \
informational and academic purposes only and does not constitute investment advice.<br>
Generated by Financial Analysis AI Agent.
</div>

</body>
</html>"""


# ============================================================
# DOCX report
# ============================================================
def build_docx_report(company_info: Dict, year: int, sections: Dict[str, str],
                       charts: List[Dict], ratios: Dict,
                       compare_df: pd.DataFrame) -> bytes:
    """Build a Word document report - returns bytes."""
    from docx import Document
    from docx.shared import Inches, Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Title
    doc.add_heading(
        f"{company_info.get('name', '')} ({company_info.get('ticker', '')})", level=0)
    subtitle = doc.add_paragraph(f"FY{year} Financial Analysis Report")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Meta info
    mcap = company_info.get("market_cap")
    mcap_str = f"${mcap/1e9:.1f}B {company_info.get('currency', 'USD')}" if mcap else "-"

    meta_p = doc.add_paragraph()
    meta_p.add_run(f"Sector: {company_info.get('sector', '-')} / "
                    f"{company_info.get('industry', '-')}\n").italic = True
    meta_p.add_run(f"Country: {company_info.get('country', '-')}  |  ").italic = True
    meta_p.add_run(f"Market Cap: {mcap_str}").italic = True

    doc.add_paragraph(
        "This report is auto-generated by the AI Financial Analysis Agent. "
        "Each section first presents key data (charts and tables), followed by "
        "AI-driven analysis built on that data."
    )

    # ===== Helper functions =====
    def add_md_paragraphs(text: str):
        """Render Markdown text into docx, handling **bold**, lists, ### headers."""
        if not text or not text.strip():
            doc.add_paragraph("(No analysis available for this section.)").italic = True
            return
        for line in text.split("\n"):
            line = line.rstrip()
            if not line.strip():
                continue
            if line.startswith("### "):
                doc.add_heading(line[4:], level=3)
            elif line.startswith("- ") or line.startswith("* "):
                p = doc.add_paragraph(style="List Bullet")
                _add_runs_with_bold(p, line[2:])
            else:
                p = doc.add_paragraph()
                _add_runs_with_bold(p, line)

    def add_ratios_table(keys: List[str]):
        """Add a table for a subset of ratios."""
        valid_rows = []
        for k in keys:
            v = ratios.get(k)
            if v is None and k not in ratios:
                continue
            disp = _format_value(v, k)
            valid_rows.append((k, disp))
        if not valid_rows:
            return
        table = doc.add_table(rows=1, cols=2)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Metric"
        hdr[1].text = "Value"
        for k, disp in valid_rows:
            row = table.add_row().cells
            row[0].text = k
            row[1].text = disp

    def add_charts_for(anchor: str):
        """Embed all charts matching the given section anchor."""
        for ch in charts:
            if ch.get("section_anchor") != anchor:
                continue
            png = fig_to_png_bytes(ch["fig"])
            if png:
                doc.add_picture(io.BytesIO(png), width=Inches(6.0))
                cap = doc.add_paragraph(f"Figure: {ch['title']}")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in cap.runs:
                    run.italic = True
                    run.font.size = Pt(9)
            else:
                p = doc.add_paragraph(f"Chart '{ch['title']}' could not be rendered")
                for run in p.runs:
                    run.italic = True

    def add_company_info_table():
        rows_data = [
            ("Company Name", company_info.get("name", "-")),
            ("Ticker", company_info.get("ticker", "-")),
            ("Sector", company_info.get("sector", "-")),
            ("Industry", company_info.get("industry", "-")),
            ("Country", company_info.get("country", "-")),
            ("Market Cap", mcap_str),
            ("Reporting Currency", company_info.get("currency", "-")),
        ]
        table = doc.add_table(rows=len(rows_data), cols=2)
        table.style = "Light Grid Accent 1"
        for i, (k, v) in enumerate(rows_data):
            table.rows[i].cells[0].text = k
            table.rows[i].cells[1].text = str(v)

    def add_peer_table():
        if compare_df.empty:
            return
        table = doc.add_table(rows=1, cols=len(compare_df.columns) + 1)
        table.style = "Light Grid Accent 1"
        hdr = table.rows[0].cells
        hdr[0].text = "Metric"
        for i, col in enumerate(compare_df.columns):
            hdr[i + 1].text = str(col)[:25]
        for idx in compare_df.index:
            row = table.add_row().cells
            row[0].text = idx
            for i, col in enumerate(compare_df.columns):
                v = compare_df.loc[idx, col]
                row[i + 1].text = _format_value(v, idx)

    # Metric groups for each section
    profit_keys = ["Gross Margin", "Operating Margin", "Net Margin",
                   "ROA (Return on Assets)", "ROE (Return on Equity)"]
    operating_keys = ["Asset Turnover", "Inventory Turnover", "Receivables Turnover"]
    solvency_keys = ["Current Ratio", "Quick Ratio", "Debt to Assets",
                     "Equity Multiplier", "Interest Coverage"]

    # 1. Company Overview
    doc.add_heading("1. Company Overview", level=1)
    add_company_info_table()
    add_md_paragraphs(sections.get("overview", ""))

    # 2. Profitability
    doc.add_heading("2. Profitability Analysis", level=1)
    add_charts_for("profitability")
    doc.add_heading("Key Profitability Metrics", level=2)
    add_ratios_table(profit_keys)
    add_md_paragraphs(sections.get("profitability", ""))

    # 3. Operating Efficiency
    doc.add_heading("3. Operating Efficiency", level=1)
    doc.add_heading("Efficiency Metrics", level=2)
    add_ratios_table(operating_keys)
    add_md_paragraphs(sections.get("operating", ""))

    # 4. Solvency
    doc.add_heading("4. Solvency and Capital Structure", level=1)
    add_charts_for("solvency")
    doc.add_heading("Solvency Metrics", level=2)
    add_ratios_table(solvency_keys)
    add_md_paragraphs(sections.get("solvency", ""))

    # 5. DuPont
    doc.add_heading("5. DuPont Analysis: ROE Drivers", level=1)
    add_charts_for("dupont")
    add_md_paragraphs(sections.get("dupont", ""))

    # 6. Trends
    doc.add_heading("6. Multi-Year Trends", level=1)
    add_charts_for("trend")
    add_md_paragraphs(sections.get("trend", ""))

    # 7. Peer Comparison
    doc.add_heading("7. Peer Comparison", level=1)
    if not compare_df.empty:
        doc.add_heading("Peer Comparison Table", level=2)
        add_peer_table()
        add_charts_for("peer")
    add_md_paragraphs(sections.get("peer", ""))

    # 8. Diagnosis
    doc.add_heading("8. Overall Diagnosis and Investor Perspective", level=1)
    add_md_paragraphs(sections.get("diagnosis", ""))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _add_runs_with_bold(paragraph, text):
    """Handle **bold** markup when rendering to docx."""
    import re
    parts = re.split(r"(\*\*[^\*]+\*\*)", text)
    for part in parts:
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


# ============================================================
# PDF report (via HTML)
# ============================================================
def build_pdf_report(html_content: str) -> Optional[bytes]:
    """
    HTML -> PDF using weasyprint. Returns None if weasyprint is not available.
    """
    try:
        from weasyprint import HTML
        return HTML(string=html_content).write_pdf()
    except Exception:
        return None
